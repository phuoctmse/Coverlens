import json
from pathlib import Path

import openpyxl
import pytest

from coverlens.adapters.markdown_spec import MarkdownSpecAdapter
from coverlens.adapters.xlsx_suite import XlsxSuiteAdapter
from coverlens.core.cascade import run_cascade
from coverlens.core.models import Case, CoverageResult, Requirement, Status, Verdict
from coverlens.core.pack import load_pack
from coverlens.core.protocols import OutputWriter
from coverlens.core.report import Report, build_report
from coverlens.verifiers.fake import FakeVerifier
from coverlens.verifiers.noop import NoopVerifier
from coverlens.writers.excel_writer import ExcelWriter
from coverlens.writers.json_writer import JsonWriter

ROOT = Path(__file__).resolve().parents[1]


def sample_report() -> Report:
    reqs = [
        Requirement(id="A.1", parent_id="A", text="First"),
        Requirement(id="A.2", parent_id="A", text="Second"),
    ]
    cases = [
        Case(id="T-1", title="t", steps="", expected_result="", refs=("A.1",)),
        Case(id="T-2", title="t", steps="", expected_result="", refs=("A.2",)),
        Case(id="T-3", title="t", steps="", expected_result=""),
    ]
    result = CoverageResult(
        verdicts=(
            Verdict(
                requirement_id="A.1",
                status=Status.COVERED,
                tier=3,
                cited_case_ids=("T-1",),
                rationale="ok",
            ),
            Verdict(requirement_id="A.2", status=Status.GAP, tier=1, rationale="none"),
        ),
        orphan_case_ids=("T-3",),
        candidate_ids={"A.1": ("T-1", "T-3"), "A.2": ()},
    )
    return build_report(reqs, cases, result, pack_name="sample")


def sheet_rows(path: Path, sheet: str) -> list[tuple[object, ...]]:
    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = list(workbook[sheet].iter_rows(values_only=True))
    workbook.close()
    return rows


@pytest.mark.parametrize("writer", [JsonWriter(), ExcelWriter()])
def test_writers_satisfy_the_protocol_and_create_the_directory(
    writer: OutputWriter, tmp_path: Path
) -> None:
    assert isinstance(writer, OutputWriter)
    path = writer.write(sample_report(), tmp_path / "new" / "dir")
    assert path.is_file()
    assert path.parent == tmp_path / "new" / "dir"


# --- JSON ----------------------------------------------------------------------


def test_json_round_trips(tmp_path: Path) -> None:
    report = sample_report()
    path = JsonWriter().write(report, tmp_path)
    assert path.name == "coverage_report.json"
    assert Report.model_validate_json(path.read_text(encoding="utf-8")) == report


def test_json_uses_plain_values(tmp_path: Path) -> None:
    data = json.loads(JsonWriter().write(sample_report(), tmp_path).read_text("utf-8"))
    assert data["requirements"][0]["status"] == "covered"
    assert data["summary"]["coverage"] == 0.5


# --- Excel ---------------------------------------------------------------------


def test_excel_has_exactly_the_two_sheets(tmp_path: Path) -> None:
    path = ExcelWriter().write(sample_report(), tmp_path)
    assert path.name == "coverage_report.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True)
    assert workbook.sheetnames == ["Gap report", "Coverage"]
    workbook.close()


def test_gap_report_has_one_row_per_requirement(tmp_path: Path) -> None:
    rows = sheet_rows(ExcelWriter().write(sample_report(), tmp_path), "Gap report")
    assert rows[0] == (
        "Requirement",
        "Story",
        "Requirement text",
        "Status",
        "Tier",
        "Cited cases",
        "Candidates",
        "Rationale",
    )
    assert rows[1] == ("A.1", "A", "First", "covered", 3, "T-1", "T-1, T-3", "ok")
    assert rows[2] == ("A.2", "A", "Second", "gap", 1, None, None, "none")
    assert len(rows) == 3


def test_coverage_sheet_has_summary_stories_and_case_findings(tmp_path: Path) -> None:
    rows = sheet_rows(ExcelWriter().write(sample_report(), tmp_path), "Coverage")
    cells = [cell for row in rows for cell in row if cell is not None]
    for expected in (
        "Summary",
        "Coverage",
        "50%",
        "Stories",
        "partial",
        "Orphan cases",
        "T-3",
        "Mis-referenced cases",
        "T-2",
        "A.2 was judged a gap: no case covers it.",
    ):
        assert expected in cells, expected


# --- the ott_web data, offline -------------------------------------------------


def test_ott_offline_report_is_written(tmp_path: Path) -> None:
    pack = load_pack(ROOT / "domains" / "ott_web" / "pack.yaml")
    data = ROOT / "data" / "ott_web"
    reqs = MarkdownSpecAdapter().read(data / "user_stories.md", pack)
    cases = XlsxSuiteAdapter().read(data / "test_cases.xlsx", pack)
    result = run_cascade(
        reqs, cases, pack, [NoopVerifier(), FakeVerifier(glossary=pack.glossary)]
    )
    report = build_report(reqs, cases, result, pack.name)
    ExcelWriter().write(report, tmp_path)
    JsonWriter().write(report, tmp_path)

    assert len(sheet_rows(tmp_path / "coverage_report.xlsx", "Gap report")) == 53
    mis_referenced = {m.case_id for m in report.mis_referenced}
    assert "TC-017" in mis_referenced  # claims US-04.AC3, which nothing covers
