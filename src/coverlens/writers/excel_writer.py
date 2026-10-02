"""Writes the report as an Excel workbook with "Gap report" and "Coverage" sheets."""

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from coverlens.core.report import Report

FILE_NAME = "coverage_report.xlsx"

BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {
    "covered": PatternFill("solid", fgColor="C6EFCE"),
    "partial": PatternFill("solid", fgColor="FFEB9C"),
    "uncertain": PatternFill("solid", fgColor="FFEB9C"),
    "gap": PatternFill("solid", fgColor="FFC7CE"),
}

GAP_COLUMNS: list[tuple[str, int]] = [
    ("Requirement", 14),
    ("Story", 10),
    ("Requirement text", 60),
    ("Status", 12),
    ("Tier", 6),
    ("Cited cases", 16),
    ("Candidates", 30),
    ("Rationale", 70),
]


def _ids(ids: Sequence[str]) -> str | None:
    return ", ".join(ids) or None


def _write_gap_report(sheet: Worksheet, report: Report) -> None:
    sheet.append([name for name, _ in GAP_COLUMNS])
    for index, (_, width) in enumerate(GAP_COLUMNS, start=1):
        sheet.cell(row=1, column=index).font = BOLD
        sheet.column_dimensions[
            sheet.cell(row=1, column=index).column_letter
        ].width = width
    for row in report.requirements:
        sheet.append(
            [
                row.requirement_id,
                row.story_id,
                row.text,
                row.status.value,
                row.tier,
                _ids(row.cited_case_ids),
                _ids(row.candidate_ids),
                row.rationale,
            ]
        )
        line = sheet.max_row
        sheet.cell(row=line, column=4).fill = STATUS_FILL[row.status.value]
        for column in (3, 7, 8):
            sheet.cell(row=line, column=column).alignment = WRAP
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _heading(sheet: Worksheet, title: str, columns: Sequence[str] = ()) -> None:
    if sheet.max_row > 1:
        sheet.append([])
    sheet.append([title])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=13)
    if columns:
        sheet.append(list(columns))
        for index in range(1, len(columns) + 1):
            sheet.cell(row=sheet.max_row, column=index).font = BOLD


def _write_coverage(sheet: Worksheet, report: Report) -> None:
    summary = report.summary
    _heading(sheet, "Summary")
    for label, value in (
        ("Pack", report.pack_name),
        ("Requirements", summary.requirements),
        ("Cases", summary.cases),
        ("Covered", summary.covered),
        ("Gap", summary.gap),
        ("Uncertain", summary.uncertain),
        ("Coverage", f"{summary.coverage:.0%}"),
    ):
        sheet.append([label, value])

    _heading(
        sheet,
        "Stories",
        ("Story", "Requirements", "Covered", "Gap", "Uncertain", "Status"),
    )
    for story in report.stories:
        sheet.append(
            [
                story.story_id,
                story.total,
                story.covered,
                story.gap,
                story.uncertain,
                story.status.value,
            ]
        )
        sheet.cell(row=sheet.max_row, column=6).fill = STATUS_FILL[story.status.value]

    _heading(sheet, "Orphan cases", ("Case",))
    for case_id in report.orphan_case_ids or ("(none)",):
        sheet.append([case_id])

    _heading(sheet, "Mis-referenced cases", ("Case", "Claimed requirement", "Reason"))
    for item in report.mis_referenced:
        sheet.append([item.case_id, item.requirement_id, item.reason])
    if not report.mis_referenced:
        sheet.append(["(none)"])

    for letter, width in zip("ABCDEF", (22, 20, 50, 10, 10, 10), strict=True):
        sheet.column_dimensions[letter].width = width


class ExcelWriter:
    def write(self, report: Report, out_dir: Path) -> Path:
        out_dir.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        gap_sheet = workbook.active
        assert gap_sheet is not None
        gap_sheet.title = "Gap report"
        _write_gap_report(gap_sheet, report)
        _write_coverage(workbook.create_sheet("Coverage"), report)
        path = out_dir / FILE_NAME
        workbook.save(path)
        return path
