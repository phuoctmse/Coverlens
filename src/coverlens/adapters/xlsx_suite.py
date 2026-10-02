"""Reads an existing test suite from an .xlsx sheet, using the pack's column map."""

import zipfile
from collections.abc import Iterable, Sequence
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from coverlens.core.models import Case
from coverlens.core.pack import PackConfig, SuiteConfig


class SuiteError(ValueError):
    """The suite is missing, unreadable or malformed."""


def cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def split_list(text: str, separator: str) -> tuple[str, ...]:
    return tuple(item for part in text.split(separator) if (item := part.strip()))


def parse_rows(
    header: Sequence[object],
    rows: Iterable[Sequence[object]],
    suite: SuiteConfig,
    source: str = "<suite>",
) -> list[Case]:
    """Turn a header and data rows into cases. Row numbers count the header as 1."""
    names = [cell_text(cell) for cell in header]
    repeated = sorted({name for name in names if name and names.count(name) > 1})
    if repeated:
        raise SuiteError(f"{source}: duplicate column names: {repeated}")
    column = {name: index for index, name in enumerate(names) if name}

    columns = suite.columns
    needed = {c for c in columns.model_dump().values() if c} | set(suite.dimensions)
    missing = sorted(needed - column.keys())
    if missing:
        raise SuiteError(f"{source}: missing columns: {missing}")

    cases: list[Case] = []
    seen: set[str] = set()
    for row_no, row in enumerate(rows, start=2):
        cells = [cell_text(value) for value in row]
        if not any(cells):
            continue
        cells += [""] * (len(names) - len(cells))

        def get(name: str | None, cells: list[str] = cells) -> str:
            return cells[column[name]] if name else ""

        where = f"{source} row {row_no}"
        case_id = get(columns.id)
        if not case_id:
            raise SuiteError(f"{where}: empty {columns.id!r}")
        if case_id in seen:
            raise SuiteError(f"{where}: duplicate case ID {case_id}")
        seen.add(case_id)
        try:
            case = Case(
                id=case_id,
                title=get(columns.title),
                preconditions=get(columns.preconditions),
                steps=get(columns.steps),
                expected_result=get(columns.expected_result),
                refs=split_list(get(columns.refs), suite.list_separator),
                tags=split_list(get(columns.tags), suite.list_separator),
                dimensions={d: v for d in suite.dimensions if (v := get(d))},
            )
        except ValidationError as exc:
            raise SuiteError(f"{where}: invalid case\n{exc}") from exc
        cases.append(case)
    if not cases:
        raise SuiteError(f"{source}: no cases found")
    return cases


class XlsxSuiteAdapter:
    """InputAdapter for an .xlsx suite."""

    def read(self, path: Path, pack: PackConfig) -> list[Case]:
        sheet = pack.suite.sheet
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
            raise SuiteError(f"{path}: cannot read suite: {exc}") from exc
        try:
            if sheet not in workbook.sheetnames:
                raise SuiteError(
                    f"{path}: no sheet {sheet!r}; found {workbook.sheetnames}"
                )
            rows = workbook[sheet].iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise SuiteError(f"{path}: sheet {sheet!r} is empty")
            return parse_rows(header, rows, pack.suite, source=str(path))
        finally:
            workbook.close()
